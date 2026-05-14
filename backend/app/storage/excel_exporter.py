from __future__ import annotations

from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


ATTENDANCE_HEADERS = [
    ("id", "记录ID"),
    ("student_no", "学号"),
    ("student_name", "姓名"),
    ("class_name", "班级"),
    ("status", "考勤状态"),
    ("capture_mode", "采集方式"),
    ("frame_count", "上传帧数"),
    ("valid_frame_count", "有效帧数"),
    ("match_score", "匹配分数"),
    ("threshold", "匹配阈值"),
    ("liveness_result", "活体结果"),
<<<<<<< HEAD
    ("liveness_mode", "活体模式"),
    ("liveness_score", "活体分数"),
=======
    ("liveness_score", "活体分数"),
    ("liveness_confidence", "活体置信度"),
    ("liveness_mode", "活体模式"),
    ("verification_mode", "验证模式"),
    ("challenge_id", "挑战ID"),
    ("challenge_action", "挑战动作"),
    ("challenge_action_label", "挑战名称"),
    ("challenge_prompt", "挑战提示"),
    ("challenge_result", "挑战结果"),
    ("challenge_score", "挑战分数"),
    ("challenge_threshold", "挑战阈值"),
    ("challenge_reason", "挑战说明"),
    ("attack_suspected", "疑似攻击"),
>>>>>>> 98bf8e49 (update)
    ("emotion", "情绪"),
    ("created_at", "考勤时间"),
    ("note", "备注"),
]

SECURITY_TEST_HEADERS = [
    ("id", "记录ID"),
    ("test_type", "测试类型"),
    ("result", "测试结果"),
    ("frame_count", "上传帧数"),
    ("valid_frame_count", "有效帧数"),
<<<<<<< HEAD
    ("score", "活体分数"),
    ("confidence", "置信度"),
    ("threshold", "阈值"),
    ("liveness_mode", "活体模式"),
=======
    ("challenge_action", "挑战动作"),
    ("challenge_action_label", "挑战名称"),
    ("challenge_prompt", "挑战提示"),
    ("challenge_result", "挑战结果"),
    ("challenge_score", "挑战分数"),
    ("challenge_confidence", "挑战置信度"),
    ("threshold", "阈值"),
    ("verification_mode", "验证模式"),
    ("challenge_id", "挑战ID"),
    ("challenge_reason", "挑战说明"),
    ("attack_suspected", "疑似攻击"),
>>>>>>> 98bf8e49 (update)
    ("operator_username", "操作人"),
    ("sample_path", "样本路径"),
    ("remark", "备注"),
    ("created_at", "测试时间"),
]

STATISTICS_DASHBOARD_HEADERS = [
    ("metric", "指标"),
    ("value", "数值"),
]

STATISTICS_FREQUENCY_HEADERS = [
    ("student_no", "学号"),
    ("student_name", "姓名"),
    ("class_name", "班级"),
    ("activity_count", "参与次数"),
]

STATISTICS_ACCURACY_HEADERS = [
    ("title", "活动名称"),
    ("activity_type", "活动类型"),
    ("activity_date", "活动日期"),
    ("participant_count", "识别参与人数"),
    ("actual_student_count", "实际参与人数"),
    ("correct_match_count", "正确识别人数"),
    ("accuracy", "准确率"),
]

STATISTICS_TYPE_HEADERS = [
    ("activity_type", "活动类型"),
    ("count", "活动数量"),
]

GROUP_PHOTO_SUMMARY_HEADERS = [
    ("activity_id", "活动ID"),
    ("title", "活动名称"),
    ("activity_type", "活动类型"),
    ("activity_date", "活动日期"),
    ("face_count", "检测人脸数"),
    ("matched_count", "匹配成功数"),
    ("unmatched_count", "未匹配数"),
    ("participant_count", "唯一参与人数"),
    ("evaluation_accuracy", "评估准确率"),
    ("created_at", "创建时间"),
]

GROUP_PHOTO_FACE_HEADERS = [
    ("id", "记录ID"),
    ("face_index", "人脸序号"),
    ("status", "状态"),
    ("student_no", "学号"),
    ("student_name", "姓名"),
    ("class_name", "班级"),
    ("match_score", "匹配分数"),
    ("det_score", "检测分数"),
    ("emotion", "情绪"),
    ("emotion_score", "情绪分数"),
    ("emotion_mode", "情绪模式"),
    ("bbox", "边框坐标"),
    ("manual_review_status", "人工修正状态"),
    ("manual_review_action", "人工修正动作"),
    ("manual_review_note", "人工修正备注"),
    ("manual_review_by", "修正人"),
    ("manual_review_at", "修正时间"),
    ("created_at", "识别时间"),
]

<<<<<<< HEAD
=======
AUDIT_LOG_HEADERS = [
    ("id", "日志ID"),
    ("created_at", "记录时间"),
    ("operator_username", "操作人"),
    ("operator_role", "操作角色"),
    ("module", "模块"),
    ("action", "动作"),
    ("result", "结果"),
    ("target_type", "目标类型"),
    ("target_id", "目标ID"),
    ("target_name", "目标对象"),
    ("message", "消息"),
    ("detail", "详情"),
]

>>>>>>> 98bf8e49 (update)

def _write_table(
    sheet,
    headers: Sequence[tuple[str, str]],
    rows: Sequence[Mapping[str, object]],
) -> None:
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx, (_, header) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = header_alignment

    for row_idx, record in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(headers, start=1):
            sheet.cell(row=row_idx, column=col_idx, value=record.get(key))

    for col_idx, (key, header) in enumerate(headers, start=1):
        max_length = len(str(header))
        for row in rows:
            value = row.get(key)
            max_length = max(max_length, len("" if value is None else str(value)))
        column_letter = sheet.cell(row=1, column=col_idx).column_letter
<<<<<<< HEAD
        sheet.column_dimensions[column_letter].width = min(max(max_length + 4, 12), 34)
=======
        sheet.column_dimensions[column_letter].width = min(max(max_length + 4, 12), 36)
>>>>>>> 98bf8e49 (update)


def export_attendance_records(
    records: Iterable[Mapping[str, object]],
    output_path: Path,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "attendance"

    rows = list(records)
    _write_table(sheet, ATTENDANCE_HEADERS, rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def export_security_test_records(
    records: Iterable[Mapping[str, object]],
    output_path: Path,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "security_tests"

    rows = list(records)
    _write_table(sheet, SECURITY_TEST_HEADERS, rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def export_statistics_report(
    report: Mapping[str, object],
    output_path: Path,
) -> Path:
    workbook = Workbook()

    dashboard = report.get("dashboard", {})
    dashboard_rows = [
        {"metric": "活动总数", "value": dashboard.get("total_activities", 0)},
<<<<<<< HEAD
        {"metric": "检测到的人脸总数", "value": dashboard.get("total_faces_detected", 0)},
        {"metric": "累计唯一参与人数", "value": dashboard.get("total_unique_participants", 0)},
        {"metric": "已评估活动数", "value": dashboard.get("evaluated_activity_count", 0)},
=======
        {
            "metric": "检测到的人脸总数",
            "value": dashboard.get("total_faces_detected", 0),
        },
        {
            "metric": "累计唯一参与人数",
            "value": dashboard.get("total_unique_participants", 0),
        },
        {
            "metric": "已评估活动数",
            "value": dashboard.get("evaluated_activity_count", 0),
        },
>>>>>>> 98bf8e49 (update)
        {"metric": "平均识别准确率", "value": dashboard.get("average_accuracy", 0)},
    ]
    dashboard_sheet = workbook.active
    dashboard_sheet.title = "dashboard"
    _write_table(dashboard_sheet, STATISTICS_DASHBOARD_HEADERS, dashboard_rows)

    frequency_sheet = workbook.create_sheet("activity_frequency")
    _write_table(
        frequency_sheet,
        STATISTICS_FREQUENCY_HEADERS,
        list(report.get("activity_frequency", [])),
    )

    accuracy_sheet = workbook.create_sheet("activity_accuracy")
    _write_table(
        accuracy_sheet,
        STATISTICS_ACCURACY_HEADERS,
        list(report.get("activity_accuracy", [])),
    )

    distribution_sheet = workbook.create_sheet("type_distribution")
    _write_table(
        distribution_sheet,
        STATISTICS_TYPE_HEADERS,
        list(report.get("activity_type_distribution", [])),
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def export_group_photo_activity(
    summary: Mapping[str, object],
    face_rows: Iterable[Mapping[str, object]],
    output_path: Path,
) -> Path:
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "activity_summary"
    _write_table(summary_sheet, GROUP_PHOTO_SUMMARY_HEADERS, [summary])

    faces_sheet = workbook.create_sheet("face_results")
    normalized_rows = []
    for row in face_rows:
        item = dict(row)
        bbox = item.get("bbox")
        if bbox is not None and not isinstance(bbox, str):
            item["bbox"] = str(bbox)
        normalized_rows.append(item)
    _write_table(faces_sheet, GROUP_PHOTO_FACE_HEADERS, normalized_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
<<<<<<< HEAD
=======


def export_audit_log_records(
    records: Iterable[Mapping[str, object]],
    output_path: Path,
) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "audit_logs"

    normalized_rows = []
    for row in records:
        item = dict(row)
        detail = item.get("detail")
        if detail is not None and not isinstance(detail, str):
            item["detail"] = str(detail)
        normalized_rows.append(item)

    _write_table(sheet, AUDIT_LOG_HEADERS, normalized_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
>>>>>>> 98bf8e49 (update)
